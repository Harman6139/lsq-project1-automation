"""Independent source, likelihood, count recursion, and workbook checks."""
from pathlib import Path
from itertools import product
import json,re,hashlib
import numpy as np
import pandas as pd
from scipy.stats import norm,binom
from pypdf import PdfReader
from openpyxl import load_workbook
from compute_assignment7_regimes import forward,viterbi,moments

ROOT=Path(__file__).resolve().parents[1]
r=json.loads((ROOT/'assignment7_results.json').read_text())
d=pd.read_csv(ROOT/'data/fund_x_2017_01_to_2026_07.csv')
provenance=json.loads((ROOT/'data/source_provenance.json').read_text())
source=ROOT/'source/LSQ_July2026_source_snapshot.pdf'
assert hashlib.sha256(source.read_bytes()).hexdigest()==provenance['snapshot_sha256']
assert hashlib.sha256((ROOT/'source/Fund_X_two_regime_model.docx').read_bytes()).hexdigest()==provenance['source_docx_sha256']
text='\n'.join(p.extract_text() for p in PdfReader(source).pages)
section=text.split('Monthly Performance*',1)[1].split('Statistics*',1)[0]
parsed={}
for line in section.splitlines():
    if re.match(r'^20\d\d ',line):
        year=int(line[:4]);values=re.findall(r'([+-]?\d+\.\d+)%([e]?)',line[4:])[:-2]
        for month,(v,flag) in enumerate(values,1):parsed[f'{year:04d}-{month:02d}']=(float(v)/100,flag)
assert len(parsed)==115
for row in d.itertuples():assert abs(parsed[row.date][0]-row.lsq_return)<1e-14
assert [x for x,v in parsed.items() if v[1]]==['2026-07']

f={k:np.array(r['best'][k]) for k in ['mu','sd','P','pi']}
assert np.allclose(f['pi']@f['P'],f['pi'])
yy=np.log1p(d.lsq_return.to_numpy())
for missing in [False,True]:
    small=yy[:6].copy()
    if missing:small[2]=np.nan
    likelihoods=[];paths=list(product([0,1],repeat=6))
    for path in paths:
        ll=f['pi'][path[0]]
        for t in range(6):
            if t:ll*=f['P'][path[t-1],path[t]]
            if np.isfinite(small[t]):ll*=norm.pdf(small[t],f['mu'][path[t]],f['sd'][path[t]])
        likelihoods.append(ll)
    actual=forward(small,*(f[k] for k in ['mu','sd','P','pi']))[0].sum()
    assert abs(actual-np.log(sum(likelihoods)))<1e-11
    if not missing:assert list(viterbi(small,f))==list(paths[np.argmax(likelihoods)])
assert abs(forward(yy,*(f[k] for k in ['mu','sd','P','pi']))[0].sum()-r['best']['ll'])<1e-9
assert abs(r['best']['ll']-r['independent_statsmodels']['ll'])<1e-5
for k in ['mu','sd','P','pi']:assert np.allclose(r['best'][k],r['independent_statsmodels'][k],atol=3e-6)
mix={k:np.array(r['iid'][k]) for k in ['mu','sd','P','pi']}
q=np.dot(mix['pi'],norm.cdf(-mix['mu']/mix['sd']))
assert abs(moments(mix,115)['count_le_5_probability']-binom.cdf(5,115,q))<1e-12
assert abs(forward(yy,*(mix[k] for k in ['mu','sd','P','pi']))[0].sum()-r['iid']['ll'])<1e-7
permutation=pd.read_csv(ROOT/'tables/ordering_permutations.csv').iloc[:,0].to_numpy()
boot=pd.read_csv(ROOT/'tables/refit_bootstrap.csv').iloc[:,0].to_numpy()
assert len(permutation)==499 and len(boot)==199 and permutation.min()>-1e-6
assert (1+(permutation>=r['ordering_test']['gain']).sum())/500==r['ordering_test']['p']
assert (1+(boot>=r['forecast_Q12_bootstrap']['Q12']).sum())/200==r['forecast_Q12_bootstrap']['p']

book=ROOT/'Assignment_7_Fund_X_Two_Regime_Review.xlsx'
wb=load_workbook(book,data_only=True); wf=load_workbook(book,data_only=False)
assert wb.sheetnames==['Results','Monthly returns','Forecast checks','Sensitivity']
for sh in wb:
    for row in sh:
        for cell in row:assert cell.data_type!='e',(sh.title,cell.coordinate,cell.value)
ws=wb['Monthly returns']
for i,row in enumerate(d.itertuples(),6):
    assert ws.cell(i,1).value.strftime('%Y-%m')==row.date
    assert abs(ws.cell(i,2).value-row.lsq_return)<1e-14
    assert abs(ws.cell(i,3).value-np.log1p(row.lsq_return))<1e-14
    assert wf['Monthly returns'].cell(i,3).value==f'=LN(1+B{i})'
summary=wb['Results']
assert summary['B28'].value==115 and summary['B29'].value==5
assert abs(summary['B30'].value-r['data']['mean'])<1e-14
assert abs(summary['B31'].value-r['data']['sd'])<1e-14
for i,model in enumerate(r['comparison'],6):
    for col,key in [('C','log_likelihood'),('D','AIC'),('E','BIC')]:assert abs(summary[f'{col}{i}'].value-model[key])<1e-10
assert abs(summary['E24'].value-.006)<1e-12 and abs(summary['E25'].value-.08)<1e-12
for i in range(2):
    assert abs(summary.cell(19+i,6).value-r['best_moments']['mean_returns'][i])<1e-12
    assert abs(summary.cell(19+i,7).value-r['best_moments']['sd_returns'][i])<1e-12
    assert abs(summary.cell(19+i,8).value-r['best_moments']['durations'][i])<1e-9
for j,key in enumerate(['one','mixture','hmm'],2):
    assert abs(wb['Forecast checks'].cell(28,j).value-r['validation'][key]['log_score'])<1e-10
    assert abs(wb['Forecast checks'].cell(29,j).value-r['validation'][key]['rmse'])<1e-12
for i,s in enumerate(r['sensitivity'],6):
    assert wb['Sensitivity'].cell(i,2).value==s['n']
    assert abs(wb['Sensitivity'].cell(i,7).value-s['moments']['durations'][1])<1e-10
pdf=PdfReader(ROOT/'Assignment_7_Fund_X_Two_Regime_Review.pdf')
assert len(pdf.pages)==3
assert '\u2014' not in (ROOT/'Assignment_7_Fund_X_Two_Regime_Review.tex').read_text()
manifest_path=ROOT/'assignment7_manifest.json';manifest=json.loads(manifest_path.read_text())
manifest['checks'].update({'independent_source_pdf_reconciliation':115,'brute_force_64_path_likelihood_and_missing_month_check':True,'viterbi_brute_force_check':True,'markov_loss_count_matches_iid_binomial_special_case':True,'monte_carlo_counts_verified':True,'xlsx_all_returns_formulas_and_cached_results_verified':True,'xlsx_formula_errors':0,'pdf_pages':3})
manifest_path.write_text(json.dumps(manifest,indent=2)+'\n')
print('PASS: 115 source returns; independent likelihood and missing-month checks; Monte Carlo counts; all workbook returns, formulas and headline values; 3-page PDF.')
